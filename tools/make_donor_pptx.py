"""Lay out a target charset in a precise grid, one PPTX slide per page, font
forced to a named typeface (e.g. a CYBBCJT-style font that refuses to embed).
No calibration marks needed -- every cell's EMU rect is known exactly at
generation time, so cropping after PDF export is pure arithmetic, not
computer vision.

Usage: python tools/make_donor_pptx.py --chars-file donor-target-charset.txt \
    --font CYBBCJT --out build/donor-CYBBCJT/donor.pptx --cols 10 --rows 10
"""
import argparse
import json
import os

from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

EMU_PER_IN = 914400


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--chars-file", required=True)
    ap.add_argument("--font", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--cols", type=int, default=10)
    ap.add_argument("--rows", type=int, default=10)
    ap.add_argument("--cell-in", type=float, default=1.0)
    ap.add_argument("--font-pt", type=float, default=54)
    args = ap.parse_args()

    with open(args.chars_file, encoding="utf-8") as fh:
        chars = [c for c in fh.read() if c.strip()]

    per_slide = args.cols * args.rows
    cell_emu = int(args.cell_in * EMU_PER_IN)
    slide_w = cell_emu * args.cols
    slide_h = cell_emu * args.rows

    prs = Presentation()
    prs.slide_width = Emu(slide_w)
    prs.slide_height = Emu(slide_h)
    blank = prs.slide_layouts[6]

    manifest = {"cell_in": args.cell_in, "cols": args.cols, "rows": args.rows,
                "slide_w_emu": slide_w, "slide_h_emu": slide_h, "pages": []}

    for i in range(0, len(chars), per_slide):
        page_chars = chars[i:i + per_slide]
        slide = prs.slides.add_slide(blank)
        page_idx = len(manifest["pages"])
        cells = []
        for j, ch in enumerate(page_chars):
            row, col = divmod(j, args.cols)
            left = col * cell_emu
            top = row * cell_emu
            box = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(cell_emu), Emu(cell_emu))
            tf = box.text_frame
            tf.word_wrap = False
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = 0
            tf.margin_right = 0
            tf.margin_top = 0
            tf.margin_bottom = 0
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run()
            run.text = ch
            run.font.size = Pt(args.font_pt)
            run.font.name = args.font
            cells.append({"row": row, "col": col, "char": ch, "cp": ord(ch)})
        manifest["pages"].append({"index": page_idx, "cells": cells})

    os.makedirs(os.path.dirname(args.out), exist_ok=True)
    prs.save(args.out)
    manifest_path = os.path.splitext(args.out)[0] + "_manifest.json"
    with open(manifest_path, "w", encoding="utf-8") as fh:
        json.dump(manifest, fh, ensure_ascii=False, indent=1)

    print(f"chars: {len(chars)}  slides: {len(manifest['pages'])}  "
          f"per_slide: {per_slide}")
    print(f"pptx: {args.out}")
    print(f"manifest: {manifest_path}")


if __name__ == "__main__":
    main()
